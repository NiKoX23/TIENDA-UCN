import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook

app = FastAPI(title="Analitica Tienda UCN")
RUTA_EXCEL = Path(
    os.environ.get(
        "INVENTARIO_XLSX_PATH",
        str(Path(__file__).resolve().parent.parent / "docs" / "Productos.xlsx"),
    )
)


def _numero(valor: Any) -> int:
    try:
        return int(valor or 0)
    except (TypeError, ValueError):
        return 0


def _estado(stock: int) -> str:
    if stock <= 5:
        return "CRITICO"
    if stock <= 20:
        return "BAJO"
    if stock <= 50:
        return "NORMAL"
    return "ALTO"


def _categoria(nombre: str) -> str:
    nombre = nombre.upper()
    if "POLERON" in nombre:
        return "Polerones"
    if "LIBRETA" in nombre or "LAPIZ" in nombre:
        return "Papeler\u00eda"
    return "Accesorios"


def _codigo_producto(sku: str) -> str:
    sku = sku.upper()
    if sku.startswith("POL-"):
        return "POLERON-TALLAS"
    if sku.startswith("LAPIZ-"):
        return "LAPIZ-COBRE"
    if sku.startswith("LIBRETA-"):
        return "LIBRETA-UCN"
    if sku.startswith("TOTE-"):
        return "TOTE-BAG-UCN"
    if sku.startswith("PELUCHE-"):
        return "PELUCHE-UCN"
    if sku.startswith("LLAVERO-UCENIN-"):
        return "LLAVERO-UCENIN"
    for prefijo, codigo in (
        ("TORTUGA-", "LL-TORTUGA"), ("ESTRELLA-", "LL-ESTRELLA"),
        ("CANGREJO-ERMITANO-", "LL-ERM-CANGREJO"), ("CANGREJO-", "LL-CANGREJO"),
        ("SOLDEMAR-", "LL-SOLDEMAR"), ("PEZ-BALLESTA-", "LL-PEZ-BALLESTA"),
        ("GAVIOTA-", "LL-GAVIOTA"), ("CACHALOTE-", "LL-CACHALOTE"),
        ("PINGUINO-", "LL-PINGUINO"), ("TIBURON-", "LL-TIBURON"),
    ):
        if sku.startswith(prefijo):
            return codigo
    return "OTRO"


NOMBRES = {
    "POLERON-TALLAS": "Poler\u00f3n UCN", "LAPIZ-COBRE": "Lapicera de cobre",
    "LIBRETA-UCN": "Libreta UCN", "TOTE-BAG-UCN": "Tote bag UCN",
    "LLAVERO-UCENIN": "Llavero UCENIN", "PELUCHE-UCN": "Peluche UCN rosa",
    "LL-TORTUGA": "Llavero Tortuga", "LL-ESTRELLA": "Llavero Estrella",
    "LL-ERM-CANGREJO": "Llavero Cangrejo ermita\u00f1o", "LL-CANGREJO": "Llavero Cangrejo",
    "LL-SOLDEMAR": "Llavero Soldemar", "LL-PEZ-BALLESTA": "Llavero Pez ballesta",
    "LL-GAVIOTA": "Llavero Gaviota", "LL-CACHALOTE": "Llavero Cachalote",
    "LL-PINGUINO": "Llavero Ping\u00fcino", "LL-TIBURON": "Llavero Tibur\u00f3n",
}


def _leer_datos(libro):
    inventario = libro["Inventario"]
    productos: dict[str, dict[str, Any]] = {}
    por_sku: dict[str, dict[str, Any]] = {}
    for fila in inventario.iter_rows(min_row=3, min_col=1, max_col=7, values_only=True):
        nombre, color, talla, precio, precio_tac, stock, sku = fila
        if not nombre or not sku:
            continue
        nombre = str(nombre).strip()
        sku = str(sku).strip()
        codigo = _codigo_producto(sku)
        variante = {
            "nombre": NOMBRES.get(codigo, nombre.title()),
            "codigoProducto": codigo,
            "categoria": _categoria(nombre),
            "color": str(color).strip().lower() if color else None,
            "talla": str(talla).strip() if talla else "unica",
            "precio": _numero(precio),
            "precioTac": _numero(precio_tac),
            "stock": _numero(stock),
            "sku": sku,
        }
        por_sku[sku] = variante
        producto = productos.setdefault(
            codigo,
            {"codigoProducto": codigo, "nombre": NOMBRES.get(codigo, nombre.title()), "categoria": _categoria(nombre),
             "stock": 0, "vendidos": 0, "ingresoNormal": 0, "ingresoTac": 0,
             "cantInicial": 0},
        )
        producto["stock"] += variante["stock"]

    ventas = []
    hoja_ventas = libro["Registro de ventas"]
    ultimo_id = None
    ultima_fecha = None
    for numero_fila, fila in enumerate(
        hoja_ventas.iter_rows(min_row=3, min_col=1, max_col=5, values_only=True), start=3
    ):
        folio, fecha, sku, cantidad, total = fila
        if not sku or not cantidad:
            continue
        sku = str(sku).strip()
        variante = por_sku.get(sku)
        if not variante:
            continue
        folio_num = None
        if isinstance(folio, (int, float)) and float(folio).is_integer():
            folio_num = int(folio)
        elif folio is None:
            folio_num = ultimo_id
        if folio_num is not None:
            ultimo_id = folio_num
        if isinstance(fecha, (datetime, date)):
            fecha_str = fecha.isoformat()
            ultima_fecha = fecha_str
        elif fecha:
            fecha_str = str(fecha)
            ultima_fecha = fecha_str
        else:
            fecha_str = ultima_fecha or ""

        unidades = _numero(cantidad)
        subtotal = _numero(total)
        ganancia = subtotal - unidades * variante["precioTac"]
        ventas.append({
            "id": numero_fila,
            "numeroVenta": folio_num,
            "fecha": fecha_str,
            "producto": variante["nombre"],
            "sku": sku,
            "variante": " ".join(filter(None, [variante["color"], variante["talla"] if variante["talla"] != "unica" else ""])),
            "cantidad": unidades,
            "metodoPago": "",
            "tipoVenta": "normal",
            "subtotal": subtotal,
            "ganancia": ganancia,
        })

    tac = []
    hoja_tac = libro["CONTROL TAC"]
    for numero_fila, fila in enumerate(
        hoja_tac.iter_rows(min_row=3, min_col=1, max_col=8, values_only=True), start=3
    ):
        n_tac, fecha, sku, cantidad, total, firma_comprador, firma_vendedor, responsable = fila
        if n_tac is None or not sku or not cantidad:
            continue
        sku = str(sku).strip()
        variante = por_sku.get(sku)
        if not variante:
            continue
        fecha_str = fecha.isoformat() if isinstance(fecha, (datetime, date)) else str(fecha or "")
        tac.append({
            "fila": numero_fila, "numero": str(n_tac).strip(), "fecha": fecha_str,
            "sku": sku, "cantidad": _numero(cantidad), "total": _numero(total),
            "firmaComprador": str(firma_comprador or "Pendiente").strip(),
            "firmaVendedor": str(firma_vendedor or "Pendiente").strip(),
            "responsable": str(responsable or "").strip(), "variante": variante,
        })
    return productos, por_sku, ventas, tac


def _abrir():
    if not RUTA_EXCEL.is_file():
        raise FileNotFoundError(f"No existe el archivo {RUTA_EXCEL}")
    return load_workbook(RUTA_EXCEL, data_only=True, read_only=True)


@app.get("/")
def estado():
    return {"servicio": "analitica", "estado": "activo"}


@app.get("/resumen")
def resumen():
    try:
        libro = _abrir()
        try:
            productos, variantes, ventas, tac = _leer_datos(libro)
        finally:
            libro.close()
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"No se pudo leer docs/Productos.xlsx: {exc}") from exc

    vendidos_por_sku: dict[str, int] = {}
    for venta in ventas:
        sku = venta["sku"]
        codigo = _codigo_producto(sku)
        producto = productos[codigo]
        producto["vendidos"] += venta["cantidad"]
        producto["ingresoNormal"] += venta["subtotal"]
        vendidos_por_sku[sku] = vendidos_por_sku.get(sku, 0) + venta["cantidad"]
    for documento in tac:
        codigo = documento["variante"]["codigoProducto"]
        producto = productos[codigo]
        producto["vendidos"] += documento["cantidad"]
        producto["ingresoTac"] += documento["total"]
        vendidos_por_sku[documento["sku"]] = vendidos_por_sku.get(documento["sku"], 0) + documento["cantidad"]

    for producto in productos.values():
        producto["cantInicial"] = producto["stock"] + producto["vendidos"]
        producto["porcentajeVendido"] = round(producto["vendidos"] / producto["cantInicial"] * 100, 1) if producto["cantInicial"] else 0

    unidades_normales = sum(v["cantidad"] for v in ventas)
    unidades_tac = sum(t["cantidad"] for t in tac)
    ingreso_normal = sum(v["subtotal"] for v in ventas)
    ingreso_tac = sum(t["total"] for t in tac)
    ganancia = sum(v["ganancia"] for v in ventas)
    ids_normales = {v["numeroVenta"] for v in ventas if v["numeroVenta"] is not None}
    ventas_sin_folio = sum(1 for v in ventas if v["numeroVenta"] is None)
    ids_tac = {t["numero"] for t in tac}
    tac_por_id = {}
    for documento in tac:
        tac_por_id.setdefault(documento["numero"], documento)
    total_stock = sum(p["stock"] for p in productos.values())
    total_vendidas = unidades_normales + unidades_tac
    total_inicial = total_stock + total_vendidas

    alertas = []
    for sku, variante in variantes.items():
        stock = variante["stock"]
        if stock <= 20:
            alertas.append({
                "codigoProducto": variante["codigoProducto"], "nombre": variante["nombre"],
                "sku": sku, "talla": variante["talla"], "color": variante["color"],
                "stock": stock, "stockMinimo": 0, "estado": _estado(stock),
            })
    resumen_lista = list(productos.values())
    resumen_lista.sort(key=lambda p: p["nombre"].lower())
    alertas.sort(key=lambda a: (a["estado"], a["stock"]))
    return {
        "general": {
            "totalProductos": len(productos), "totalUnidadesIniciales": total_inicial,
            "stockActualTotal": total_stock, "unidadesVendidas": total_vendidas,
            "porcentajeVendido": round(total_vendidas / total_inicial * 100, 1) if total_inicial else 0,
        },
        "ventas": {
            "totalVentas": len(ids_normales) + ventas_sin_folio + len(ids_tac),
            "ventasNormales": len(ids_normales) + ventas_sin_folio, "ventasTac": len(ids_tac),
            "ingresoNormal": ingreso_normal, "ingresoTac": ingreso_tac,
            "ingresoTotal": ingreso_normal + ingreso_tac, "gananciaTotal": ganancia,
        },
        "controlTac": {
            "totalDocumentos": len(ids_tac),
            "aprobados": sum(1 for t in tac_por_id.values() if t["firmaComprador"].casefold() == "firmado" and t["firmaVendedor"].casefold() == "firmado"),
            "pendientes": len(ids_tac) - sum(1 for t in tac_por_id.values() if t["firmaComprador"].casefold() == "firmado" and t["firmaVendedor"].casefold() == "firmado"),
            "montoTotal": ingreso_tac,
        },
        "resumenPorProducto": resumen_lista, "alertasStockBajo": alertas,
    }


@app.get("/registro-ventas")
def registro_ventas():
    try:
        libro = _abrir()
        try:
            _, _, ventas, _ = _leer_datos(libro)
        finally:
            libro.close()
        return ventas
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"No se pudo leer el registro de ventas: {exc}") from exc
